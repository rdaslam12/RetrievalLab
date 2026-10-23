from io import BytesIO

import fitz
from docx import Document

from src.document_parser import DocumentValidationError, parse_corpus, parse_document


def test_empty_extracted_pdf_pages_produce_warnings() -> None:
    document = fitz.open()
    document.new_page()
    page = document.new_page()
    page.insert_text((72, 72), "Extractable scientific text")
    payload = document.tobytes()
    document.close()

    sections, warnings = parse_document("paper.pdf", payload)
    assert [section.page_number for section in sections] == [2]
    assert any("page 1" in warning and "no extractable text" in warning for warning in warnings)


def test_docx_preserves_paragraph_order_without_page_number() -> None:
    document = Document()
    document.add_paragraph("First paragraph.")
    document.add_paragraph("Second paragraph.")
    buffer = BytesIO()
    document.save(buffer)

    sections, warnings = parse_document("notes.docx", buffer.getvalue())
    assert warnings == []
    assert sections[0].page_number is None
    assert sections[0].text == "First paragraph.\nSecond paragraph."


def test_duplicate_filenames_receive_distinct_source_ids() -> None:
    parsed = parse_corpus([("same.txt", b"one"), ("same.txt", b"two")])
    assert parsed.files_processed == 2
    assert len({section.source_id for section in parsed.sections}) == 2


def test_txt_decode_fallback_is_reported() -> None:
    sections, warnings = parse_document("legacy.txt", b"valid\xfftext")
    assert sections and "\ufffd" in sections[0].text
    assert any("invalid UTF-8" in warning for warning in warnings)


def test_beir_jsonl_creates_one_identified_section_per_document() -> None:
    payload = (
        b'{"_id":"doc-1","title":"First title","text":"First body","metadata":{"x":1}}\n'
        b'\n'
        b'{"_id":2,"title":"","text":"Second body"}\n'
    )

    sections, warnings = parse_document("corpus.jsonl", payload)

    assert warnings == []
    assert [section.document_id for section in sections] == ["doc-1", "2"]
    assert [section.title for section in sections] == ["First title", None]
    assert [section.text for section in sections] == ["First title\n\nFirst body", "Second body"]
    assert all(section.file_type == "jsonl" for section in sections)
    assert all(section.file_name == "corpus.jsonl" for section in sections)
    assert all(section.page_number is None for section in sections)
    assert len({section.source_id for section in sections}) == 2


def test_jsonl_empty_documents_are_skipped_with_a_warning() -> None:
    payload = (
        b'{"_id":"empty","title":"","text":""}\n'
        b'{"_id":"kept","text":"Indexable body"}\n'
    )

    sections, warnings = parse_document("corpus.jsonl", payload)

    assert [section.document_id for section in sections] == ["kept"]
    assert any("'empty'" in warning and "was skipped" in warning for warning in warnings)


def test_jsonl_reports_malformed_line_number() -> None:
    try:
        parse_document(
            "corpus.jsonl",
            b'{"_id":"valid","text":"Body"}\n{"_id":"broken",}\n',
        )
    except DocumentValidationError as exc:
        assert "line 2" in str(exc)
        assert "not valid JSON" in str(exc)
    else:
        raise AssertionError("Expected malformed JSONL to fail")


def test_jsonl_rejects_duplicate_document_ids() -> None:
    payload = b'{"_id":"same","text":"One"}\n{"_id":"same","text":"Two"}\n'
    try:
        parse_document("corpus.jsonl", payload)
    except DocumentValidationError as exc:
        assert "duplicate document _id" in str(exc)
        assert "line 2" in str(exc)
    else:
        raise AssertionError("Expected duplicate JSONL document IDs to fail")


def test_jsonl_requires_beir_id_and_text_fields() -> None:
    invalid_payloads = [
        (b'{"text":"Body"}\n', "'_id'"),
        (b'{"_id":"doc"}\n', "'text'"),
    ]
    for payload, expected in invalid_payloads:
        try:
            parse_document("corpus.jsonl", payload)
        except DocumentValidationError as exc:
            assert expected in str(exc)
        else:
            raise AssertionError("Expected incomplete BEIR JSONL record to fail")
